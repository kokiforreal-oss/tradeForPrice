from __future__ import annotations

import csv
from io import BytesIO, StringIO
from typing import Any, Iterable, Optional

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

EXPORT_HEADERS = [
    ("sku", "产品ID"),
    ("name", "产品名称"),
    ("spec", "规格"),
    ("category_code", "产品分类编码"),
    ("category_name", "产品分类"),
    ("product_type", "商品类型"),
    ("pricing_method", "计价方式"),
    ("unit", "计量单位"),
    ("primary_unit", "主计量单位"),
    ("aux_unit", "辅助计量单位"),
    ("sales_unit", "销售单位"),
    ("cost_price", "参考成本"),
    ("status", "状态"),
    ("remark", "备注"),
]

HEADER_ALIASES = {
    "sku": ("产品id", "sku", "产品编号", "货号"),
    "name": ("产品名称", "名称", "name", "品名"),
    "spec": ("规格", "型号", "spec"),
    "category_code": ("产品分类编码", "分类编码", "categorycode"),
    "category_name": ("产品分类", "分类名称", "分类", "categoryname", "category"),
    "product_type": ("商品类型", "产品类型", "producttype"),
    "pricing_method": ("计价方式", "pricingmethod"),
    "unit": ("计量单位", "单位", "unit"),
    "primary_unit": ("主计量单位", "primaryunit"),
    "aux_unit": ("辅助计量单位", "auxunit"),
    "sales_unit": ("销售单位", "salesunit"),
    "cost_price": ("参考成本", "成本", "cost", "costprice"),
    "status": ("状态", "status"),
    "remark": ("备注", "remark", "说明"),
}

MAX_IMPORT_BYTES = 8 * 1024 * 1024
MAX_IMPORT_ROWS = 3000

_HEADER_FILL = PatternFill("solid", fgColor="2563EB")
_HEADER_FONT = Font(color="FFFFFF", bold=True)


def cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "是" if value else "否"
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return format(value, "g")
    if isinstance(value, int):
        return str(value)
    return str(value).strip().replace("\ufeff", "")


def _norm_header(value: Any) -> str:
    return cell_text(value).lower().replace(" ", "").replace("_", "").replace("　", "")


def _header_index(headers: list[str]) -> dict[str, int]:
    alias_to_field = {}
    for field, aliases in HEADER_ALIASES.items():
        for a in aliases:
            alias_to_field[_norm_header(a)] = field
    mapping: dict[str, int] = {}
    for i, h in enumerate(headers):
        field = alias_to_field.get(_norm_header(h))
        if field and field not in mapping:
            mapping[field] = i
    return mapping


def _row_to_item(headers_map: dict[str, int], cells: list[Any], row_no: int) -> Optional[dict]:
    if not any(cell_text(c) for c in cells):
        return None
    item = {"_row": row_no}
    for field, idx in headers_map.items():
        item[field] = cell_text(cells[idx] if idx < len(cells) else "")
    return item


def parse_csv(raw: bytes) -> list[dict]:
    text = None
    for enc in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise ValueError("无法识别 CSV 编码，请另存为 UTF-8")
    reader = csv.reader(StringIO(text))
    rows = list(reader)
    if not rows:
        raise ValueError("文件是空的")
    mapping = _header_index(rows[0])
    if "sku" not in mapping or "name" not in mapping:
        raise ValueError("缺少必填列：产品ID、产品名称")
    out = []
    for i, row in enumerate(rows[1:], start=2):
        item = _row_to_item(mapping, row, i)
        if item:
            out.append(item)
        if len(out) > MAX_IMPORT_ROWS:
            raise ValueError(f"最多导入 {MAX_IMPORT_ROWS} 行")
    return out


def parse_xlsx(raw: bytes) -> list[dict]:
    wb = load_workbook(BytesIO(raw), read_only=True, data_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        rows_iter = ws.iter_rows(values_only=True)
        header_row = next(rows_iter, None)
        if not header_row:
            raise ValueError("文件是空的")
        mapping = _header_index([cell_text(c) for c in header_row])
        if "sku" not in mapping or "name" not in mapping:
            raise ValueError("缺少必填列：产品ID、产品名称")
        out = []
        for i, row in enumerate(rows_iter, start=2):
            item = _row_to_item(mapping, list(row), i)
            if item:
                out.append(item)
            if len(out) > MAX_IMPORT_ROWS:
                raise ValueError(f"最多导入 {MAX_IMPORT_ROWS} 行")
        return out
    finally:
        wb.close()


def parse_upload(filename: str, raw: bytes) -> list[dict]:
    name = (filename or "").lower()
    if name.endswith(".xls") and not name.endswith(".xlsx"):
        raise ValueError("暂不支持旧版 .xls，请另存为 .xlsx 或 .csv")
    if name.endswith(".xlsx"):
        return parse_xlsx(raw)
    if name.endswith(".csv") or name.endswith(".txt"):
        return parse_csv(raw)
    if raw[:2] == b"PK":
        return parse_xlsx(raw)
    return parse_csv(raw)


def _write_xlsx(rows: Iterable[dict], include_cost: bool) -> bytes:
    headers = [h for k, h in EXPORT_HEADERS if include_cost or k != "cost_price"]
    keys = [k for k, _ in EXPORT_HEADERS if include_cost or k != "cost_price"]
    wb = Workbook()
    ws = wb.active
    ws.title = "产品"
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
    for row in rows:
        ws.append([row.get(k, "") for k in keys])
    for i, h in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(22, max(12, len(h) + 4))
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_template(include_cost: bool) -> bytes:
    sample = {
        "sku": "DEMO-001",
        "name": "示例产品（导入前请删除本行）",
        "spec": "规格A",
        "category_code": "",
        "category_name": "",
        "product_type": "实物",
        "pricing_method": "固定价",
        "unit": "pcs",
        "primary_unit": "pcs",
        "aux_unit": "",
        "sales_unit": "pcs",
        "cost_price": "0",
        "status": "启用",
        "remark": "",
    }
    return _write_xlsx([sample], include_cost)


def build_xlsx(rows: Iterable[dict], include_cost: bool) -> bytes:
    return _write_xlsx(rows, include_cost)


def build_csv(rows: Iterable[dict], include_cost: bool) -> bytes:
    headers = [h for k, h in EXPORT_HEADERS if include_cost or k != "cost_price"]
    keys = [k for k, _ in EXPORT_HEADERS if include_cost or k != "cost_price"]
    buf = StringIO()
    writer = csv.writer(buf)
    writer.writerow(headers)
    for row in rows:
        writer.writerow([row.get(k, "") for k in keys])
    return buf.getvalue().encode("utf-8-sig")
